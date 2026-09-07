"""
gerar_dicionario.py
--------------------
Gera docs/dicionario_dados.xlsx a partir das tabelas processadas (data/processed/*.csv).
Uso: python scripts/gerar_dicionario.py  (rode depois do tratamento_dados.py)
"""

import pandas as pd
from pathlib import Path

PROCESSED_DIR = Path("data/processed")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PROCESSED_DIR = Path("data/processed")
DOCS_DIR = Path("docs")
DOCS_DIR.mkdir(exist_ok=True)

DESCRICOES = {
    "ean": "Código de barras universal do produto (chave de junção entre todas as tabelas).",
    "produto": "Descrição comercial do produto.",
    "marca": "Marca fabricante do produto.",
    "fornecedor": "Fornecedor responsável pelo produto.",
    "categoria": "Categoria comercial do produto (ex.: Similares, Genéricos, Não Medicamentos).",
    "secao": "Seção/departamento de prateleira do produto.",
    "classe_terapeutica": "Classe terapêutica à qual o produto pertence.",
    "competencia": "Mês/ano de referência da venda (formato AAAA-MM-01).",
    "unidades": "Quantidade vendida no mês (unidades).",
    "valor": "Faturamento do mês para o produto (R$).",
    "flag_devolucao": "Verdadeiro quando 'unidades' é negativo (indício de devolução ou estorno).",
    "flag_outlier_valor": "Verdadeiro quando 'valor' está acima do percentil 99, que mostra possível erro de lançamento.",
    "cobertura_dias": "Quantos dias o estoque atual cobre, considerando a venda média recente.",
    "cobertura_semanas": "Mesma métrica de cobertura, convertida para semanas.",
    "classificacao_estoque": "Categoria de risco do estoque: Risco de Ruptura / Baixo / Ideal / Atenção / Excesso.",
    "nivel_servico": "Percentual de disponibilidade do produto no período (1 = 100%).",
    "estoque_disponivel_und": "Quantidade física disponível para venda imediata (unidades).",
    "total_estoque_rs": "Valor financeiro total em estoque (R$), incluindo transferências e balanceamentos.",
    "excesso_rs": "Valor em R$ acima da cobertura ideal definida para a curva do produto.",
    "media_venda_3m_und": "Média de unidades vendidas nos últimos 3 meses fechados.",
    "market_share": "Participação de mercado estimada do produto (0 a 1).",
    "unidades_totais": "Soma de unidades vendidas em todo o histórico (jan/2024–jul/2026).",
    "valor_total": "Soma do faturamento em todo o histórico (R$).",
    "media_mensal_unidades": "Média de unidades vendidas por mês, considerando todo o histórico.",
    "media_mensal_valor": "Média de faturamento mensal, considerando todo o histórico.",
    "meses_com_venda": "Quantidade de meses, dos 31 disponíveis, em que houve venda > 0.",
    "crescimento_trimestral": "Variação percentual entre a média dos últimos 3 meses e os 3 meses anteriores.",
    "flag_divergencia_estoque": "Verdadeiro quando a soma dos componentes de estoque não bate com o Total Estoque informado.",
    "total_estoque_recalculado": "Total de estoque recalculado para conferência com a fonte.",
}



def descrever_coluna(nome: str) -> str:
    return DESCRICOES.get(nome, "Coluna herdada da base de origem — descrição a validar com o time de negócio.")


def montar_dicionario(df: pd.DataFrame, nome_tabela: str) -> pd.DataFrame:
    linhas = []
    for col in df.columns:
        serie = df[col]
        exemplo = serie.dropna().iloc[0] if serie.notna().any() else ""
        linhas.append({
            "tabela": nome_tabela,
            "coluna": col,
            "tipo_dado": str(serie.dtype),
            "pct_nulos": round(100 * serie.isna().mean(), 1),
            "exemplo": str(exemplo)[:60],
            "descricao": descrever_coluna(col),
        })
    return pd.DataFrame(linhas)


def exportar_xlsx(dicionarios: dict[str, pd.DataFrame], caminho: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True, name="Arial")
    body_font = Font(name="Arial", size=10)

    for nome_aba, df in dicionarios.items():
        ws = wb.create_sheet(nome_aba[:31])
        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for _, row in df.iterrows():
            ws.append(list(row))

        for i, col in enumerate(df.columns, start=1):
            largura = max(14, min(60, df[col].astype(str).str.len().max() + 2))
            ws.column_dimensions[get_column_letter(i)].width = largura
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center", wrap_text=(cell.column_letter == "F"))

    wb.save(caminho)


def main_dicionario():
    tabelas = {
        "dim_produto": pd.read_csv(PROCESSED_DIR / "dim_produto.csv", dtype={"ean": str}),
        "fato_vendas": pd.read_csv(PROCESSED_DIR / "fato_vendas.csv", dtype={"ean": str}),
        "fato_estoque": pd.read_csv(PROCESSED_DIR / "fato_estoque.csv", dtype={"ean": str}),
        "kpi_produto": pd.read_csv(PROCESSED_DIR / "kpi_produto.csv", dtype={"ean": str}),
    }
    dicionarios = {nome: montar_dicionario(df, nome) for nome, df in tabelas.items()}
    saida = DOCS_DIR / "dicionario_dados.xlsx"
    exportar_xlsx(dicionarios, saida)
    print("Dicionário salvo em:", saida.resolve())
    for nome, d in dicionarios.items():
        print(f"  {nome}: {len(d)} colunas documentadas")

main_dicionario()

